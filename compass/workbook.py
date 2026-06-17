from __future__ import annotations

import csv
import math
from pathlib import Path
from typing import Dict, List, Sequence

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


MEASUREMENT_COLUMNS = [
    "comet_area",
    "comet_length",
    "comet_dna_content",
    "comet_average_intensity",
    "head_area",
    "head_diameter",
    "head_dna_content",
    "head_average_intensity",
    "head_dna_percent",
    "tail_area",
    "tail_length",
    "tail_dna_content",
    "tail_dna_percent",
    "tail_moment",
    "olive_moment",
]

DETECTION_COLUMNS = [
    "image_name",
    "relative_path",
    "image_path",
    "treatment",
    "detection_index",
    "good_bad",
    "detection_source",
    "review_action",
    "auto_selected",
    "score",
    "damage_class",
    "label",
    "overlap",
    "cutoff",
    "empty",
    "box_x1",
    "box_y1",
    "box_x2",
    "box_y2",
    *MEASUREMENT_COLUMNS,
]

SUMMARY_COLUMNS = [
    "image_name",
    "relative_path",
    "image_path",
    "treatment",
    "total_detections",
    "selected_good",
    "rejected_bad",
    "manual_added",
    "manual_deselected",
    "mean_tail_dna_percent",
    "mean_tail_moment",
    "mean_olive_moment",
]

PERCENT_COLUMNS = {"head_dna_percent", "tail_dna_percent", "mean_tail_dna_percent"}


def finite_or_none(value: object) -> object:
    if isinstance(value, (float, np.floating)) and not math.isfinite(float(value)):
        return None
    return value


def write_measurements_xlsx(workbook_path: Path, rows: Sequence[Dict[str, object]]) -> None:
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "measurements"

    good_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    bad_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
    deselected_fill = PatternFill(fill_type="solid", fgColor="F4B183")
    manual_added_fill = PatternFill(fill_type="solid", fgColor="9DC3E6")
    header_font = Font(bold=True)

    ws.append(["Selected" if col == "good_bad" else col for col in DETECTION_COLUMNS])
    for cell in ws[1]:
        cell.font = header_font
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append([
            selected_cell_value(row) if col == "good_bad" else finite_or_none(row.get(col))
            for col in DETECTION_COLUMNS
        ])
        excel_row = ws.max_row
        status_cell = ws.cell(row=excel_row, column=DETECTION_COLUMNS.index("good_bad") + 1)
        if row.get("review_action") == "manual_deselected":
            status_cell.fill = deselected_fill
        elif row.get("review_action") == "manual_added" or row.get("detection_source") == "manual":
            status_cell.fill = manual_added_fill
        else:
            status_cell.fill = good_fill if bool(status_cell.value) else bad_fill
        for col_name in PERCENT_COLUMNS.intersection(DETECTION_COLUMNS):
            col_idx = DETECTION_COLUMNS.index(col_name) + 1
            ws.cell(row=excel_row, column=col_idx).number_format = "0.00%"

    for idx, col_name in enumerate(DETECTION_COLUMNS, start=1):
        max_width = max(len(str(col_name)), 10)
        for values in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, values_only=True):
            for value in values:
                if value is not None:
                    max_width = max(max_width, min(len(str(value)), 30))
        ws.column_dimensions[get_column_letter(idx)].width = min(max_width + 2, 32)

    wb.save(workbook_path)


def write_csv(path: Path, rows: Sequence[Dict[str, object]], columns: Sequence[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(columns))
        writer.writeheader()
        for row in rows:
            writer.writerow({col: finite_or_none(row.get(col)) for col in columns})


def selected_cell_value(row: Dict[str, object]) -> object:
    if row.get("review_action") == "manual_deselected":
        return "Deselected"
    return row.get("good_bad") == "Good"


def summarize_image_rows(image_name: str, rows: List[Dict[str, object]]) -> Dict[str, object]:
    first = rows[0] if rows else {}
    good = [row for row in rows if row.get("good_bad") == "Good"]
    manual_added = [row for row in rows if row.get("review_action") == "manual_added" or row.get("detection_source") == "manual"]
    manual_deselected = [row for row in rows if row.get("review_action") == "manual_deselected"]

    def mean_of(col: str) -> object:
        vals = [float(row[col]) for row in good if row.get(col) is not None and math.isfinite(float(row[col]))]
        return float(sum(vals) / len(vals)) if vals else None

    return {
        "image_name": image_name,
        "relative_path": first.get("relative_path", ""),
        "image_path": first.get("image_path", ""),
        "treatment": first.get("treatment", ""),
        "total_detections": int(len(rows)),
        "selected_good": int(len(good)),
        "rejected_bad": int(len(rows) - len(good)),
        "manual_added": int(len(manual_added)),
        "manual_deselected": int(len(manual_deselected)),
        "mean_tail_dna_percent": mean_of("tail_dna_percent"),
        "mean_tail_moment": mean_of("tail_moment"),
        "mean_olive_moment": mean_of("olive_moment"),
    }
